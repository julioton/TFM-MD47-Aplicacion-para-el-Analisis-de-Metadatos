import os
import time
from stat import *
from tkinter import *
from tkinter import filedialog

import PyPDF2
import exifread
import olefile
import docx
import openpyxl

from PIL import Image
from PyPDF2 import PdfFileWriter, PdfFileReader
from PyPDF2.generic import NameObject, createStringObject

from core.definitions import tdPressAnalyze
from lib import pdf_metadata


class MainMethods:

    @staticmethod
    def browse_file():
        file_types = [
            ("all files", "*.*")
        ]
        filename = filedialog.askopenfilename(initialdir="/",
                                              title="Select a File",
                                              filetypes=file_types)
        return filename

    @staticmethod
    def check_file_ext(filename):
        if filename.endswith('.doc') | filename.endswith('.docx') | filename.endswith('.xls') \
                | filename.endswith('.xlsx') | filename.endswith('.docm') | filename.endswith('.dotx') \
                | filename.endswith('.dotm') | filename.endswith('.DOC') | filename.endswith('.DOCX') \
                | filename.endswith('.XLS') | filename.endswith('.XLSX') | filename.endswith('.DOCM') \
                | filename.endswith('.DOTX') | filename.endswith('.DOTM') | filename.endswith('.PDF') \
                | filename.endswith('.PNG') | filename.endswith('.pdf') | filename.endswith('.png') \
                | filename.endswith('.ppt') | filename.endswith('.pptx') | filename.endswith('.pps') \
                | filename.endswith('.xlsm') | filename.endswith('.gif') | filename.endswith('.JPG') \
                | filename.endswith('.jpg') | filename.endswith('.jpeg') | filename.endswith('.txt') \
                | filename.endswith('.PDF'):
            return True
        else:
            return False

    @staticmethod
    def print_results(output_item, input_data):
        for key, value in input_data.items():
            output_item.insert(INSERT, str(key) + ': ' + str(value) + '\n')

    @staticmethod
    def get_file_properties(filename):
        file_prop = {}
        info = []
        try:
            info = os.stat(filename)
            file_prop["File Name"] = os.path.splitext(filename)[0]
            file_prop["Extension"] = os.path.splitext(filename)[1]
            file_prop["Size"] = info[ST_SIZE]
            file_prop["Time Created"] = time.asctime(time.localtime(info[ST_CTIME]))
            file_prop["Time Last Access"] = time.asctime(time.localtime(info[ST_ATIME]))
            file_prop["Time Modified"] = time.asctime(time.localtime(info[ST_MTIME]))
        except Exception as e:
            print("Failed ", str(filename), "Exception: ", str(e))
        return file_prop

    @staticmethod
    def get_metadata(filename, fileprops):
        if filename.endswith('.png') | filename.endswith('.gif') | filename.endswith('.JPG') \
                | filename.endswith('.jpg') | filename.endswith('.jpeg'):
            return MainMethods.get_image_metadata(filename)
        elif filename.endswith('.pdf') | filename.endswith('.PDF'):
            return MainMethods.get_pdf_metadata(filename, fileprops)
        elif filename.endswith('.XLS') | filename.endswith('.xls') | \
                filename.endswith('.doc') | filename.endswith('.DOC') \
                | filename.endswith('.docm'):
            return MainMethods.get_office_old_metadata(filename, fileprops)
        elif filename.endswith('.docx') | filename.endswith('.DOCX'):
            return MainMethods.get_office_docx_metadata(filename, fileprops)
        elif filename.endswith('.xlsx') | filename.endswith('.XLSX'):
            return MainMethods.get_office_xlsx_metadata(filename, fileprops)
        else:
            return False

    @staticmethod
    def get_image_metadata(filename):
        file_prop = {}
        result = {}
        try:
            '''
            exif_data = Image.open(filename)._getexif()
            if exif_data:
                print("Exif data")
                for key, value in exif_data.items():
                    file_prop.update({TAGS.get(key): value})
                    #print('%s = %s ' % (TAGS.get(key), value))
            '''
            f = open(filename, 'rb')
            res = exifread.process_file(f)
            if res:
                print("Exif read data")
                file_prop.update({tdPressAnalyze + 'Metadata Type ': 'EXIF\n'})
                for tag in res.keys():
                    if tag not in ('JPEGThumbnail', 'TIFFThumbnail', 'Filename', 'EXIF MakerNote'):
                        file_prop.update({tag: res[tag]})
                        # print("%s = %s" % (tag, res[tag]))
            if file_prop:
                for key, value in file_prop.items():
                    if value not in result.values():
                        result[key] = value

            if int(len(result.items())) == 0:
                result.update({'\nMetadata Type EXIF': 'No Metadata to show'})

        except Exception as e:
            print("Failed ", str(filename), "Exception: ", str(e))
        return result

    @staticmethod
    def get_pdf_metadata(filename, fileprops):
        result_py_pdf = {}
        result_py_pdf_xmp = {}
        result_pdf_meta = {}
        result = {}
        result.update({tdPressAnalyze + 'Metadata Type ': 'PDF\n'})

        try:
            filepdf1 = PyPDF2.PdfFileReader(open(filename, 'rb'))
            result_py_pdf = filepdf1.getDocumentInfo()
        except Exception as e:
            print("Failed PyPDF2 getDocumentInfo", str(filename), "Exception: ", str(e))

        try:
            filepdf2 = PyPDF2.PdfFileReader(open(filename, 'rb'))
            result_py_pdf_xmp = filepdf2.getXmpMetadata()
        except Exception as e:
            print("Failed PyPDF2 getXmpMetadata", str(filename), "Exception: ", str(e))

        try:
            arg_list = [filename, '', '']
            result_pdf_meta = pdf_metadata.input_main(arg_list)
        except Exception as e:
            print("Failed pdf_meta", str(filename), "Exception: ", str(e))

        if result_py_pdf:
            count = 0
            try:
                for key, value in result_py_pdf.items():
                    ++count
                    if value not in result.values():
                        result[key] = value
            except Exception as e:
                print("Failed for result_py_pdf", str(filename), "Exception: ", str(e))

        if result_py_pdf_xmp:
            print(result_py_pdf_xmp)
            count = 0
            try:
                for key, value in result_py_pdf_xmp.items():
                    ++count
                    if value not in result.values():
                        result[key] = value
            except Exception as e:
                print("Failed for result_py_pdf_xmp", str(filename), "Exception: ", str(e))

        if result_pdf_meta:
            count = 0
            try:
                for key, value in result_pdf_meta.items():
                    ++count
                    if value not in result.values():
                        result[key] = value
            except Exception as e:
                print("Failed for result_pdf_meta", str(filename), "Exception: ", str(e))

        if int(len(result.items())) == 0:
            result.update({'\nMetadata Type PDF': 'No Metadata to show'})

        return result

    @staticmethod
    def get_office_old_metadata(filename, fileprops):
        file_prop = {}
        result_metadata = {}
        result = {}
        try:
            f = open(filename, 'rb')
            ole = olefile.OleFileIO(filename)
            result_metadata = ole.get_metadata()


            if result_metadata:
                count = 0
                try:
                    for prop in result_metadata.SUMMARY_ATTRIBS:
                        value = getattr(result_metadata, prop)
                        ++count
                        if value not in result.values():
                            result[prop] = value
                except Exception as e:
                    print("Failed for result_metadata SUMMARY_ATTRIBS", str(filename), "Exception: ", str(e))

                try:
                    for prop in result_metadata.DOCSUM_ATTRIBS:
                        value = getattr(result_metadata, prop)
                        ++count
                        if value not in result.values():
                            result[prop] = value

                except Exception as e:
                    print("Failed for result_metadata DOCSUM_ATTRIBS", str(filename), "Exception: ", str(e))

            if int(len(result.items())) == 0:
                result.update({'\nMetadata Type MS Office': 'No Metadata to show'})

            return result

        except Exception as e:
            print("Failed ", str(filename), "Exception: ", str(e))
            result.update({'\nFailed': str(e)})
        return result

    @staticmethod
    def get_office_docx_metadata(filename, fileprops):
        file_prop = {}
        result_metadata = {}
        core_properties = {}
        result = {}
        try:
            document = docx.Document(docx=filename)
            core_properties = document.core_properties
            if core_properties:
                try:
                    result_metadata.update({'author':core_properties.author})
                    result_metadata.update({'created':core_properties.created})
                    result_metadata.update({'last_modified_by':core_properties.last_modified_by})
                    result_metadata.update({'last_printed':core_properties.last_printed})
                    result_metadata.update({'modified':core_properties.modified})
                    result_metadata.update({'revision':core_properties.revision})
                    result_metadata.update({'title':core_properties.title})
                    result_metadata.update({'category':core_properties.category})
                    result_metadata.update({'comments':core_properties.comments})
                    result_metadata.update({'identifier':core_properties.identifier})
                    result_metadata.update({'keywords':core_properties.keywords})
                    result_metadata.update({'language':core_properties.language})
                    result_metadata.update({'subject':core_properties.subject})
                    result_metadata.update({'version':core_properties.version})
                    result_metadata.update({'content_status':core_properties.content_status})
                except Exception as e:
                    print("Failed for result_metadata", str(filename), "Exception: ", str(e))
                count = 0
                for key, value in result_metadata.items():
                    ++count
                    if value not in result.values() and value != '':
                        result[key] = value
            if int(len(result.items())) == 0:
                result.update({'\nMetadata Type DOCX': 'No Metadata to show'})

            return result

        except Exception as e:
            print("Failed ", str(filename), "Exception: ", str(e))
            result.update({'\nFailed': str(e)})
        return result

    @staticmethod
    def get_office_xlsx_metadata(filename, fileprops):
        file_prop = {}
        result_metadata = {}
        core_properties = {}
        result = {}
        try:
            fh = openpyxl.load_workbook(filename)
            core_properties = fh.properties  # To get old properties
            if core_properties:
                try:
                    result_metadata.update({'creator': core_properties.creator})
                    result_metadata.update({'title': core_properties.title})
                    result_metadata.update({'description': core_properties.description})
                    result_metadata.update({'subject': core_properties.subject})
                    result_metadata.update({'identifier': core_properties.identifier})
                    result_metadata.update({'language': core_properties.language})
                    result_metadata.update({'created': core_properties.created})
                    result_metadata.update({'modified': core_properties.modified})
                    result_metadata.update({'lastModifiedBy': core_properties.lastModifiedBy})
                    result_metadata.update({'category': core_properties.category})
                    result_metadata.update({'contentStatus': core_properties.contentStatus})
                    result_metadata.update({'version': core_properties.version})
                    result_metadata.update({'revision': core_properties.revision})
                    result_metadata.update({'keywords': core_properties.keywords})
                    result_metadata.update({'lastPrinted': core_properties.lastPrinted})

                except Exception as e:
                    print("Failed for result_metadata", str(filename), "Exception: ", str(e))
                count = 0
                for key, value in result_metadata.items():
                    ++count
                    if value not in result.values() and value != '':
                        result[key] = value
            if int(len(result.items())) == 0:
                result.update({'\nMetadata Type XLSX': 'No Metadata to show'})

            return result

        except Exception as e:
            print("Failed ", str(filename), "Exception: ", str(e))
            result.update({'\nFailed': str(e)})
        return result

    @staticmethod
    def clean_metadata(simplefilename, extension):
        filename = str(simplefilename) + str(extension)
        if filename.endswith('.png') | filename.endswith('.gif') | filename.endswith('.JPG') \
                | filename.endswith('.jpg') | filename.endswith('.jpeg'):
            return MainMethods.clean_image_metadata(simplefilename, extension)
        elif filename.endswith('.pdf') | filename.endswith('.PDF'):
            return MainMethods.clean_pdf_metadata(simplefilename, extension)
        elif filename.endswith('.docx') | filename.endswith('.DOCX'):
            return MainMethods.clean_office_docx_metadata(simplefilename, extension)
        elif filename.endswith('.xlsx') | filename.endswith('.XLSX'):
            return MainMethods.clean_office_xlsx_metadata(simplefilename, extension)
        elif filename.endswith('.XLS') | filename.endswith('.xls') \
                | filename.endswith('.doc') | filename.endswith('.DOC') \
                | filename.endswith('.docm'):
            return MainMethods.clean_office_metadata(simplefilename, extension)
        else:
            return False

    @staticmethod
    def clean_image_metadata(simplefilename, extension):
        filename = str(simplefilename) + str(extension)
        try:
            image = Image.open(filename)
            data = list(image.getdata())
            image_without_exif = Image.new(image.mode, image.size)
            image_without_exif.putdata(data)
            image_without_exif.save(str(simplefilename) + '_clean_' + str(extension))
            return str(simplefilename) + '_clean_' + str(extension)
        except Exception as e:
            print("Failed ", str(filename), "Exception: ", str(e))
        return False

    @staticmethod
    def clean_pdf_metadata(simplefilename, extension):
        filename = str(simplefilename) + str(extension)
        try:

            OUTPUT = simplefilename + str('_CLEAN_.pdf')
            INPUTS = [filename, ]
            inputs = {}
            output = PdfFileWriter()

            infoDict = output._info.getObject()
            infoDict.update({
                NameObject('/Title'): createStringObject(u'Title'),
                NameObject('/Producer'): createStringObject(u'MD47'),
                NameObject('/Author'): createStringObject(u'MD47'),
                NameObject('/Subject'): createStringObject(u'subject'),
                NameObject('/Creator'): createStringObject(u'MD47')
            })

            try:
                inputs = [PdfFileReader(open(i, "rb")) for i in INPUTS]
            except Exception as e:
                print("Failed PdfFileReader", str(filename), "Exception: ", str(e))

            if len(inputs) > 0:
                try:
                    for input in inputs:
                        for page in range(input.getNumPages()):
                            output.addPage(input.getPage(page))
                except Exception as e:
                    print("Failed for input", str(filename), "Exception: ", str(e))

                try:
                    outputStream = open(OUTPUT, 'wb')
                    output.write(outputStream)
                    outputStream.close()
                except Exception as e:
                    print("Failed outputStream", str(filename), "Exception: ", str(e))
            else:
                return False

            return OUTPUT

        except Exception as e:
            print("Failed ", str(filename), "Exception: ", str(e))
        return False

    @staticmethod
    def clean_office_metadata(simplefilename, extension):
        filename = str(simplefilename) + str(extension)
        try:
            f = open(filename, 'rb')
            ole = olefile.OleFileIO(filename, write_mode=True)
            data = ole.openstream('WordDocument').read()
            '''data = data.replace(b'author', b'MD47')
            ole.write_sect(0x17, b'MD47')'''
            ole.write_stream('WordDocument', data)
            ole.close()
            print("Failed - It's not possible to change metadata of old OLE files, Please save this with new Office App version.")
            return False
            '''str(simplefilename) + '_clean_failed_' + str(extension)'''
        except Exception as e:
            print("Failed ", str(filename), "Exception: ", str(e))
        return False

    @staticmethod
    def clean_office_docx_metadata(simplefilename, extension):
        filename = str(simplefilename) + str(extension)
        try:
            document = docx.Document(docx=filename)
            core_properties = document.core_properties
            meta_fields = ["author", "category", "comments", "content_status", "identifier", "keywords",
                           "language", "subject", "title", "version", "last_modified_by"]
            for meta_field in meta_fields:
                setattr(core_properties, meta_field, "MD47")
            document.save(str(simplefilename) + '_clean_' + str(extension))

            return str(simplefilename) + '_clean_' + str(extension)
        except Exception as e:
            print("Failed ", str(filename), "Exception: ", str(e))
        return False

    @staticmethod
    def clean_office_xlsx_metadata(simplefilename, extension):
        filename = str(simplefilename) + str(extension)
        try:
            document = openpyxl.load_workbook(filename)
            document.properties.creator = "MD47"
            document.properties.title = "MD47"
            document.properties.description = "MD47"
            document.properties.subject = "MD47"
            document.properties.lastModifiedBy = "MD47"
            document.properties.category = "MD47"
            document.properties.keywords = "MD47"

            document.save(str(simplefilename) + '_clean_' + str(extension))

            return str(simplefilename) + '_clean_' + str(extension)
        except Exception as e:
            print("Failed ", str(filename), "Exception: ", str(e))
        return False
